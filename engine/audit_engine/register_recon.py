"""Register import (PRD 10.2) and reconciliation to the system of record (PRD 10.3, D-01)."""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections import defaultdict

from .findings import adjustment
from .gstin import normalise as norm_gstin
from .money import rupees, to_paise
from .repo import register_imports, register_recon, register_rows
from .snapshot import d

EXPECTED = ["Bill no.", "Date", "Job No.", "Customer", "Ledger Account", "State Code", "GST No.", "Total Cost",
            "Discount", "Packing", "Freight", "Insurance", "IGST", "CGST", "SGST", "TCS Collection", "Round",
            "Gross Amt.", "Tax", "BU", "E Invoice"]

ALIASES = {
    "bill_no": ["bill no", "bill number", "invoice no", "invoice number", "inv no"],
    "date": ["date", "invoice date", "bill date"],
    "job_no": ["job no", "job number", "order no", "so no"],
    "customer": ["customer", "customer name", "party", "party name"],
    "ledger_account": ["ledger account", "ledger", "sales ledger"],
    "state_code": ["state code", "pos", "place of supply"],
    "gstin": ["gst no", "gstin", "gst number", "gstin no"],
    "taxable": ["total cost", "taxable value", "taxable amount", "taxable"],
    "discount": ["discount"],
    "packing": ["packing"],
    "freight": ["freight"],
    "insurance": ["insurance"],
    "igst": ["igst"],
    "cgst": ["cgst"],
    "sgst": ["sgst", "utgst", "sgst utgst"],
    "tcs": ["tcs collection", "tcs"],
    "round": ["round", "round off", "rounding"],
    "gross": ["gross amt", "gross amount", "gross", "invoice value", "total"],
    "tax": ["tax", "tax rate"],
    "bu": ["bu", "business unit"],
    "e_invoice": ["e invoice", "einvoice", "e invoice status"],
}
CANON_OF_EXPECTED = {"Bill no.": "bill_no", "Date": "date", "Job No.": "job_no", "Customer": "customer",
                     "Ledger Account": "ledger_account", "State Code": "state_code", "GST No.": "gstin",
                     "Total Cost": "taxable", "Discount": "discount", "Packing": "packing", "Freight": "freight",
                     "Insurance": "insurance", "IGST": "igst", "CGST": "cgst", "SGST": "sgst",
                     "TCS Collection": "tcs", "Round": "round", "Gross Amt.": "gross", "Tax": "tax", "BU": "bu",
                     "E Invoice": "e_invoice"}
AMOUNT_FIELDS = ["taxable", "discount", "packing", "freight", "insurance", "igst", "cgst", "sgst", "tcs", "round",
                 "gross"]
FOOTING = ["taxable", "packing", "freight", "insurance", "igst", "cgst", "sgst", "tcs", "round"]


def norm_header(h) -> str:
    s = str(h or "").lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(s.split())


def norm_no(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", "", str(v)).upper()


def norm_name(v) -> str:
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower())


# Raw cells are stored as JSON exactly as read, with tags for date and time types (RI-2).

def enc(v):
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    if isinstance(v, dt.datetime):
        return {"$dt": v.isoformat()}
    if isinstance(v, dt.date):
        return {"$d": v.isoformat()}
    if isinstance(v, dt.time):
        return {"$t": v.isoformat()}
    return str(v)


def dec(v):
    if isinstance(v, dict):
        if "$dt" in v:
            return dt.datetime.fromisoformat(v["$dt"])
        if "$d" in v:
            return dt.date.fromisoformat(v["$d"])
        if "$t" in v:
            return dt.time.fromisoformat(v["$t"])
    return v


def read_table(data: bytes, file_name: str) -> tuple[str, list[list]]:
    """Return sheet name and every row as a list of raw cell values (RI-5)."""
    name = (file_name or "").lower()
    if name.endswith(".csv") or not data.startswith(b"PK"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
        return "csv", [list(r) for r in csv.reader(io.StringIO(text))]
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    for cand in wb.worksheets:
        if "sales" in cand.title.lower() or "register" in cand.title.lower():
            ws = cand
            break
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    while rows and all(c is None or c == "" for c in rows[-1]):
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    while width and all(len(r) < width or r[width - 1] in (None, "") for r in rows):
        width -= 1
    rows = [r[:width] for r in rows]
    title = ws.title
    wb.close()
    return title, rows


def find_header(rows: list[list]) -> tuple[int, dict[int, str], list[str]]:
    """Locate the header row and map column index to canonical field (RI-3)."""
    best = (-1, {}, [])
    for idx, row in enumerate(rows[:15]):
        mapping = {}
        for ci, cell in enumerate(row):
            h = norm_header(cell)
            for canon, names in ALIASES.items():
                if h in names and canon not in mapping.values():
                    mapping[ci] = canon
                    break
        if len(mapping) > len(best[1]):
            best = (idx, mapping, [str(c) if c is not None else "" for c in row])
    return best


def parse_date(v) -> str | None:
    v = dec(v)
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d-%m-%y", "%d/%m/%y",
                "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_row(raw: list, colmap: dict[int, str]) -> dict:
    out: dict = {}
    errors = []
    for ci, canon in colmap.items():
        v = dec(raw[ci]) if ci < len(raw) else None
        if canon == "bill_no":
            out[canon] = norm_no(v)
        elif canon == "date":
            out[canon] = parse_date(v)
            if v not in (None, "") and out[canon] is None:
                errors.append(f"date {v!r} not readable")
        elif canon in AMOUNT_FIELDS:
            p = to_paise(v) if v not in (None, "") else (0 if v in (None, "") else None)
            out[canon] = p
            if p is None:
                errors.append(f"{canon} {v!r} not a number")
        elif canon == "gstin":
            out[canon] = norm_gstin(v) if v not in (None, "") else None
        else:
            out[canon] = None if v is None else str(v).strip()
    blank = all(c in (None, "") for c in raw)
    text = " ".join(str(c) for c in raw if isinstance(c, str)).lower()
    if blank:
        kind = "blank"
    elif not out.get("bill_no") and "total" in text:
        kind = "total"
    else:
        kind = "data"
    out["_kind"] = kind
    out["_errors"] = errors
    return out


def import_register(ctx) -> dict:
    eng, snap, repo = ctx.eng, ctx.snap, ctx.engine.repo
    upload = next((u for u in snap.register_uploads if u["_seq"] == eng["register_upload_seq"]), None)
    if upload is None:
        from .engine import EngineError

        raise EngineError("MISSING_UPLOAD", "The register upload is not in the snapshot.", "Open a new engagement.")
    data = ctx.engine.files.get(upload["file_hash"])  # RI-1: read by hash, never modified
    sheet, rows = read_table(data, upload.get("file_name") or "")
    header_idx, colmap, header_names = find_header(rows)
    present = set(colmap.values())
    missing = [c for c in EXPECTED if CANON_OF_EXPECTED[c] not in present]
    stored = []
    for i, raw in enumerate(rows):
        enc_row = [enc(c) for c in raw]
        if i < header_idx:
            parsed = {"_kind": "preamble", "_errors": []}
        elif i == header_idx:
            parsed = {"_kind": "header", "_errors": []}
        else:
            parsed = parse_row(enc_row, colmap)
        stored.append({"engagement_id": ctx.eid, "row_no": i + 1, "raw": enc_row, "parsed": parsed})
    repo.insert(register_rows, stored)
    repo.insert(register_imports, {
        "engagement_id": ctx.eid, "file_hash": upload["file_hash"], "file_name": upload.get("file_name"),
        "upload_seq": upload["_seq"],
        "headers": {"sheet": sheet, "header_row": header_idx + 1, "names": header_names},
        "column_map": {str(k): v for k, v in colmap.items()}, "missing_columns": missing,
        "row_count": sum(1 for s in stored if s["parsed"]["_kind"] == "data"), "at": eng["opened_at"]})
    findings = []
    if missing:
        findings.append(ctx.finding(
            category="REGISTER_MISMATCH", subtype="MISSING_COLUMN", subject_type="upload",
            subject_id=upload["file_hash"], title=f"Register is missing {len(missing)} expected column(s)",
            event_seqs=[upload["_seq"]], document_ids=[upload.get("file_name")],
            facts={"missing": ", ".join(missing)}))
        ctx.engine.save_findings(findings)
    return {"file_hash": upload["file_hash"], "file_name": upload.get("file_name"), "sheet": sheet,
            "header_row": header_idx + 1, "row_count": sum(1 for s in stored if s["parsed"]["_kind"] == "data"),
            "columns": header_names, "missing_columns": missing, "findings_raised": len(findings)}


def _cmp_fields(p: dict, inv: dict) -> list[dict]:
    diffs = []

    def add(field, reg, sys, money=False):
        diffs.append({"field": field, "register": rupees(reg) if money and reg is not None else reg,
                      "system": rupees(sys) if money and sys is not None else sys,
                      "delta": (reg or 0) - (sys or 0) if money else None})

    if "date" in p and p["date"] != inv.get("invoice_date"):
        add("Date", p["date"], inv.get("invoice_date"))
    if "customer" in p and norm_name(p["customer"]) != norm_name(inv.get("customer_name")):
        add("Customer", p["customer"], inv.get("customer_name"))
    if "gstin" in p and (p["gstin"] or None) != (norm_gstin(inv.get("customer_gstin")) or None):
        add("GST No.", p["gstin"], inv.get("customer_gstin"))
    for canon, key, label in (("taxable", "taxable_value", "Total Cost"), ("igst", "igst", "IGST"),
                              ("cgst", "cgst", "CGST"), ("sgst", "sgst", "SGST"), ("gross", "gross", "Gross Amt.")):
        if canon in p and p[canon] is not None and abs(p[canon] - (inv.get(key) or 0)) > 1:
            add(label, p[canon], inv.get(key) or 0, money=True)
    return diffs


def reconcile_register(ctx) -> dict:
    repo, snap = ctx.engine.repo, ctx.snap
    imp = repo.first(register_imports, engagement_id=ctx.eid)
    colmap = set((imp["column_map"] or {}).values())
    rows = repo.select(register_rows, order_by=register_rows.c.row_no, engagement_id=ctx.eid)
    data_rows = [(r["row_no"], r["parsed"]) for r in rows if r["parsed"]["_kind"] == "data"]
    status: dict[int, list[str]] = defaultdict(list)
    refs: dict[int, list[str]] = defaultdict(list)
    findings: list[dict] = []

    def raise_(f: dict, row_nos: list[int], label: str) -> None:
        findings.append(f)
        for rn in row_nos:
            status[rn].append(label)
            refs[rn].append(f["finding_id"])

    sys_by_no: dict[str, list[dict]] = defaultdict(list)
    for inv in snap.invoices:
        if ctx.in_scope(inv):
            sys_by_no[norm_no(inv.get("invoice_no"))].append(inv)
    reg_by_no: dict[str, list[tuple[int, dict]]] = defaultdict(list)
    for rn, p in data_rows:
        reg_by_no[p.get("bill_no") or f"<blank row {rn}>"].append((rn, p))

    paired: set[str] = set()  # aggregate ids matched to a row
    for no, regs in reg_by_no.items():
        invs = sys_by_no.get(no, [])
        active = [i for i in invs if not i["_reversed"]]
        reversed_ = [i for i in invs if i["_reversed"]]
        if not invs:
            for rn, p in regs:
                f = ctx.finding(category="REGISTER_MISMATCH", subtype="NOT_IN_SYSTEM", subject_type="register_row",
                                subject_id=f"{no}" if len(regs) == 1 else f"{no}#row{rn}",
                                subject_label=f"Register row {rn}", title=f"Register bill {no} is not in the system",
                                amount_impact=p.get("taxable") or 0, period_effect="None",
                                document_ids=[no], facts={"row_no": rn, "bill_no": no, "taxable": p.get("taxable") or 0})
                raise_(f, [rn], "Not in system")
            if len(regs) > 1:
                f = _dup_finding(ctx, no, regs, 0)
                raise_(f, [rn for rn, _ in regs], "Duplicate")
            continue
        remaining = list(active)
        matched_rows = []
        for rn, p in regs:
            if not remaining:
                break
            best = min(remaining, key=lambda i: (len(_cmp_fields(p, i)), i["_seq"]))
            remaining.remove(best)
            matched_rows.append((rn, p, best))
        extra = [(rn, p) for rn, p in regs if rn not in {m[0] for m in matched_rows}]
        for rn, p, inv in matched_rows:
            paired.add(inv["_aggregate_id"])
            diffs = _cmp_fields(p, inv)
            idate = d(inv.get("invoice_date"))
            if idate and not ctx.in_period(idate):
                diffs.append({"field": "Period", "register": "listed in period", "system": f"dated {idate}",
                              "delta": None})
            if diffs:
                impact = max([abs(x["delta"] or 0) for x in diffs] + [0])
                text = "; ".join(f"{x['field']} register {x['register']} vs system {x['system']}" for x in diffs)
                f = ctx.finding(category="REGISTER_MISMATCH", subtype="FIELD_DIFFERENCE", subject_type="invoice",
                                subject_id=inv["invoice_no"], subject_label=inv.get("customer_name"),
                                title=f"Register differs from invoice {inv['invoice_no']} in "
                                      f"{', '.join(x['field'] for x in diffs)}",
                                amount_impact=impact, event_seqs=[inv["_seq"]], document_ids=[inv["invoice_no"]],
                                facts={"row_no": rn, "invoice_no": inv["invoice_no"], "fields": diffs,
                                       "fields_text": text})
                raise_(f, [rn], "Field difference")
            else:
                status[rn].append("Matched")
        if extra:
            if reversed_:
                inv = reversed_[0]
                rev = inv["_reversed"]
                f = ctx.finding(category="REGISTER_MISMATCH", subtype="REVERSED_IN_SYSTEM", subject_type="invoice",
                                subject_id=inv["invoice_no"], subject_label=inv.get("customer_name"),
                                title=f"Reversed invoice {inv['invoice_no']} is still in the register",
                                amount_impact=inv.get("taxable_value") or 0, period_effect="None",
                                event_seqs=[inv["_seq"], rev["_seq"]], document_ids=[inv["invoice_no"]],
                                facts={"invoice_no": inv["invoice_no"], "reversed_on": rev["_business_date"],
                                       "reversal_seq": rev["_seq"]})
                raise_(f, [rn for rn, _ in extra[:len(reversed_)]], "Reversed in system")
                extra = extra[len(reversed_):]
            if extra:
                f = _dup_finding(ctx, no, regs, len(invs))
                raise_(f, [rn for rn, _ in extra], "Duplicate")

    for no, invs in sys_by_no.items():
        for inv in invs:
            if inv["_reversed"] or inv["_aggregate_id"] in paired:
                continue
            if not ctx.in_period(inv.get("invoice_date")):
                continue
            f = ctx.finding(category="REGISTER_MISMATCH", subtype="MISSING_IN_REGISTER", subject_type="invoice",
                            subject_id=inv["invoice_no"] if len(invs) == 1 else f"{inv['invoice_no']}#{inv['_aggregate_id']}",
                            subject_label=inv.get("customer_name"),
                            title=f"Invoice {inv['invoice_no']} is missing from the register",
                            amount_impact=inv.get("taxable_value") or 0, event_seqs=[inv["_seq"]],
                            document_ids=[inv["invoice_no"]],
                            facts={"invoice_no": inv["invoice_no"], "invoice_date": inv.get("invoice_date"),
                                   "taxable": inv.get("taxable_value") or 0})
            findings.append(f)

    if {"gross"} <= colmap and "taxable" in colmap:
        present = [c for c in FOOTING if c in colmap]
        for rn, p in data_rows:
            if any(p.get(c) is None for c in present + ["gross"]):
                continue
            computed = sum(p[c] for c in present)
            if abs(computed - p["gross"]) > len(present):
                f = ctx.finding(category="REGISTER_MISMATCH", subtype="ROW_ARITHMETIC", subject_type="register_row",
                                subject_id=f"{p.get('bill_no')}#row{rn}", subject_label=p.get("bill_no"),
                                title=f"Register row {rn} does not foot", amount_impact=abs(computed - p["gross"]),
                                document_ids=[p.get("bill_no")],
                                facts={"row_no": rn, "bill_no": p.get("bill_no"), "computed": computed,
                                       "gross": p["gross"], "columns": present})
                raise_(f, [rn], "Arithmetic")

    reg_total = sum(p.get("taxable") or 0 for _, p in data_rows)
    ledger_total = 0
    for po in snap.postings:
        if po["account"] == "4000" and ctx.in_period(po["business_date"]) and (
                not ctx.entities or po["entity_id"] in ctx.entities):
            ledger_total += po["credit"] - po["debit"]
    if reg_total != ledger_total:
        diff = reg_total - ledger_total
        f = ctx.finding(category="REGISTER_MISMATCH", subtype="TOTAL_NOT_AGREED_TO_LEDGER", subject_type="period",
                        subject_id=f"{ctx.start}..{ctx.end}", title="Register total does not agree to Sales in the ledger",
                        amount_impact=abs(diff), period_effect="None", document_ids=[imp["file_hash"]],
                        facts={"register_total": reg_total, "ledger_total": ledger_total, "difference": diff})
        findings.append(f)

    for rr in rows:
        if rr["parsed"]["_kind"] == "total" and "taxable" in colmap:
            stated = rr["parsed"].get("taxable")
            if stated is not None and abs(stated - reg_total) > 100:
                f = ctx.finding(category="REGISTER_MISMATCH", subtype="TOTAL_ROW_MISMATCH",
                                subject_type="register_row", subject_id=f"total#row{rr['row_no']}",
                                title=f"Register total row {rr['row_no']} does not agree to its rows",
                                amount_impact=abs(stated - reg_total),
                                facts={"row_no": rr["row_no"], "stated": stated, "computed": reg_total,
                                       "column": "Total Cost"})
                raise_(f, [rr["row_no"]], "Total row differs")

    ctx.engine.save_findings(findings)
    recon_rows = []
    for rr in rows:
        rn, kind = rr["row_no"], rr["parsed"]["_kind"]
        if kind in ("preamble", "header"):
            continue
        label = "; ".join(dict.fromkeys(status.get(rn) or [])) or {"blank": "Blank row", "total": "Total row"}.get(
            kind, "Matched")
        recon_rows.append({"engagement_id": ctx.eid, "row_no": rn, "status": label, "finding_refs": refs.get(rn, [])})
    repo.insert(register_recon, recon_rows)
    by_sub: dict[str, int] = defaultdict(int)
    for f in findings:
        by_sub[f["subtype"]] += 1
    return {"rows_compared": len(data_rows), "rows_matched": sum(1 for r in recon_rows if r["status"] == "Matched"),
            "register_total": reg_total, "ledger_total": ledger_total, "findings_raised": len(findings),
            "by_subtype": dict(by_sub)}


def _dup_finding(ctx, no: str, regs: list, system_count: int) -> dict:
    extra = len(regs) - max(system_count, 1)
    taxable = regs[0][1].get("taxable") or 0
    return ctx.finding(category="REGISTER_MISMATCH", subtype="DUPLICATE_ROW", subject_type="invoice", subject_id=no,
                       title=f"Bill {no} is listed {len(regs)} times in the register",
                       amount_impact=taxable * max(extra, 1), document_ids=[no],
                       facts={"bill_no": no, "count": len(regs), "rows": ", ".join(str(r) for r, _ in regs),
                              "system_count": system_count})
