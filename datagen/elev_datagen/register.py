"""Client sales register in the layout of Fig. 1 (SD-6), and the related-party list."""

from __future__ import annotations

import csv
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Font

from .world import FY_END, FY_START, Sale, World

COLUMNS = ["Bill no.", "Date", "Job No.", "Customer", "Ledger Account", "State Code", "GST No.", "Total Cost",
           "Discount", "Packing", "Freight", "Insurance", "IGST", "CGST", "SGST", "TCS Collection", "Round",
           "Gross Amt.", "Tax", "BU", "E Invoice"]


def _r(paise: int) -> float:
    return round(paise / 100.0, 2)


def row_for(s: Sale) -> dict:
    inv = s.invoice
    return {
        "Bill no.": inv["invoice_no"],
        "Date": dt.datetime.combine(s.invoice_date, dt.time()),
        "Job No.": inv["order_id"] or "",
        "Customer": inv["customer_name"],
        "Ledger Account": "Sales Revenue",
        "State Code": inv["place_of_supply"],
        "GST No.": inv["customer_gstin"] or "",
        "Total Cost": _r(inv["taxable_value"]),
        "Discount": _r(inv["discount"]),
        "Packing": _r(inv["packing"]),
        "Freight": _r(inv["freight"]),
        "Insurance": _r(inv["insurance"]),
        "IGST": _r(inv["igst"]),
        "CGST": _r(inv["cgst"]),
        "SGST": _r(inv["sgst"]),
        "TCS Collection": _r(inv["tcs"]),
        "Round": _r(inv["round_off"]),
        "Gross Amt.": _r(inv["gross"]),
        "Tax": "GST 18%",
        "BU": inv["bu"],
        "E Invoice": "Generated" if inv.get("irn") else "Not generated",
    }


def build_rows(w: World) -> list[dict]:
    plants = w.register_plants
    keep_reversed = plants.get("keep_reversed", False)
    drop = {id(s) for s in plants.get("drop", [])}
    sales = [s for s in w.sales if s.entity == "E1" and FY_START <= s.invoice_date <= FY_END
             and (keep_reversed or not s.reversed_on) and id(s) not in drop]
    sales.sort(key=lambda s: (s.invoice_no, s.sale_id))
    rows = []
    dup = {id(s) for s in plants.get("dup", [])}
    field_changes = {}
    for s, col, val in plants.get("field", []):
        field_changes.setdefault(id(s), []).append((col, val))
    arith = {}
    for s, col, delta in plants.get("arith", []):
        arith.setdefault(id(s), []).append((col, delta))
    for s in sales:
        row = row_for(s)
        for col, val in field_changes.get(id(s), []):
            row[col] = val
        for col, delta in arith.get(id(s), []):
            row[col] = round(row[col] + delta, 2)
        rows.append(row)
        if id(s) in dup:
            rows.append(dict(row))
    for fake in plants.get("fake", []):
        rows.append(fake)
    return rows


def register_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Register"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(c) for c in COLUMNS])
        ws.cell(ws.max_row, 2).number_format = "DD-MM-YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return _normalise_zip(buf.getvalue())


def _normalise_zip(data: bytes) -> bytes:
    """Rewrite the xlsx with fixed timestamps so the same rows always give the same file hash."""
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re

                content = re.sub(rb"<dcterms:(created|modified)([^>]*)>[^<]*<",
                                 rb"<dcterms:\1\2>2026-07-10T12:00:00Z<", content)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 7, 10, 18, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, content)
    return out.getvalue()


def related_party_csv(parties: list[dict]) -> bytes:
    buf = io.StringIO()
    wr = csv.writer(buf, lineterminator="\n")
    wr.writerow(["Name", "PAN", "GSTIN", "Relationship", "Effective From"])
    for p in parties:
        wr.writerow([p["name"], p.get("pan") or "", p.get("gstin") or "", p["relationship"], p["effective_from"]])
    return buf.getvalue().encode("utf-8")
