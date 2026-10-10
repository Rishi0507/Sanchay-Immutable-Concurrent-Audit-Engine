from audit_engine.repo import register_recon, register_rows

from conftest import recall


def test_clean_register_has_no_findings(clean_engine):
    r = clean_engine.step("reconcile_register")
    assert r["findings_raised"] == 0
    assert r["register_total"] == r["ledger_total"]


def test_planted_register_recall(planted_engine, answer_key):
    planted_engine.step("reconcile_register")
    eid = planted_engine.resolve(None)["engagement_id"]
    findings = [f for f in planted_engine.repo.findings(eid) if f["detector_id"] == "D-01"]
    missed = recall(findings, [i for i in answer_key if i["detector"] == "D-01"])
    assert not missed, missed


def test_register_copied_exactly(planted_engine, datasets):
    planted_engine.step("import_register")
    eid = planted_engine.resolve(None)["engagement_id"]
    rows = planted_engine.repo.select(register_rows, order_by=register_rows.c.row_no, engagement_id=eid)
    from openpyxl import load_workbook

    wb = load_workbook(datasets / "planted" / "uploads" / "sales_register_FY2025-26.xlsx", read_only=True)
    source = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    assert len(rows) == len(source)
    from audit_engine.register_recon import dec

    for stored, src in zip(rows, source):
        assert [dec(c) for c in stored["raw"]] == src
    statuses = planted_engine.repo.select(register_recon, engagement_id=eid)
    assert any(s["status"] == "Duplicate" for s in statuses)


def test_reconcile_is_idempotent(planted_engine):
    eid = planted_engine.resolve(None)["engagement_id"]
    before = len(planted_engine.repo.findings(eid))
    again = planted_engine.step("reconcile_register")
    assert again.get("already_done")
    assert len(planted_engine.repo.findings(eid)) == before
